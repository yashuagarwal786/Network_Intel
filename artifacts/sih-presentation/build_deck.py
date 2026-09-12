from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

ROOT=Path(r'Z:\XLab\New\Network Intel')
OUT=ROOT/'artifacts/sih-presentation'
p=Presentation(r'C:\Users\tanis\Downloads\SIH2026-IDEA-Presentation-Format.pptx')
sid=p.slides._sldIdLst[-1]; p.part.drop_rel(sid.rId); p.slides._sldIdLst.remove(sid)
BLUE='204F89'; BLACK='182333'
def box(sl,x,y,w,h,text,size=20,bold=False,color=BLACK):
    s=sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=s.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.03)
    tf.margin_top=tf.margin_bottom=Inches(.02)
    for i,line in enumerate(text.split('\n')):
        q=tf.paragraphs[0] if i==0 else tf.add_paragraph(); q.text=line
        q.font.name='Arial'; q.font.size=Pt(size); q.font.bold=bold; q.font.color.rgb=RGBColor.from_string(color)
        q.space_after=Pt(4)
    return s
def remove(sl,id):
    s=next(s for s in sl.shapes if s.shape_id==id); s._element.getparent().remove(s._element)
def pic(sl,name,x,y,w):
    return sl.shapes.add_picture(str(ROOT/'artifacts/first-round/screenshots'/name), Inches(x), Inches(y), width=Inches(w))
def head(sl,y,text): return box(sl,.65,y,12,.4,text,20,True,BLUE)
notes=[]
def note(sl,purpose,script,takeaway,question,answer):
    txt=f'Purpose: {purpose}\n\nSpeaker script: {script}\n\nJudge takeaway: {takeaway}\n\nLikely question: {question}\nAnswer: {answer}'
    sl.notes_slide.notes_text_frame.text=txt; notes.append(txt)

s=p.slides[0]
remove(s,10)
box(s,.45,2.7,6.75,4.45,'Problem Statement ID: [TO FILL]\nProblem Statement Title: [TO FILL]\nTheme: [TO FILL]\nPS Category: Software\nTeam ID: [TO FILL]\nTeam Name: [REGISTERED NAME]',19,True)
box(s,.55,2.13,6.8,.46,'VEIL: Evidence-first investigation workspace',21,True,BLUE)
note(s,'Introduce the project and official submission identity.','VEIL connects fragmented investigation-style records into a reviewable evidence graph. Our prototype links reports, call records, transactions and vehicles, then lets investigators inspect the source behind each lead. The demonstration uses synthetic data. Official problem and team details must be filled before submission.','A working evidence workflow with human decision authority.','Is this deployed with an agency?','No. This is a local prototype using synthetic records. Agency integration and validation remain future work.')

s=p.slides[1]; remove(s,15362)
head(s,1.42,'Proposed Solution (Describe your Idea/Solution/Prototype)')
box(s,.65,1.98,12,.58,'VEIL connects separate records into entities, relationships and explainable leads.',23,True)
box(s,.65,2.82,5.3,.65,'Detailed explanation of the proposed solution',17,True,BLUE)
box(s,.65,3.53,5.1,.85,'Reports, CDR, transactions and vehicle records in one workspace.',20)
box(s,.65,4.38,5.3,.4,'How it addresses the problem',18,True,BLUE)
box(s,.65,4.83,5.1,.88,'Investigators can follow a connection and open its exact source record.',21)
pic(s,'03-network-path.png',6.05,2.77,6.6)
box(s,6.05,6.5,6.6,.3,'Actual prototype: a computed, sourced association path. Synthetic data.',12)
box(s,.65,5.86,5.3,.35,'Innovation and uniqueness of the solution',18,True,BLUE)
box(s,.65,6.25,5.1,.55,'Cross-source graph + provenance + reversible human review.',17)
note(s,'Explain the workflow and its differentiation.','The challenge is connecting information across separate records. VEIL builds a graph with traceable source evidence. This actual prototype screenshot shows a computed association path. Unlike manual cross-referencing, the same workspace links the relationship to its evidence and preserves reversible identity decisions. A connection is a lead for review, not proof of wrongdoing.','The differentiator is continuity from record to graph to reviewed lead.','What is novel beyond a graph viewer?','The backend combines source provenance, reviewed identity resolution, bounded traversal and evidence-gated signals within one investigator workflow.')

s=p.slides[2]; remove(s,17410)
head(s,1.42,'Technologies to be used')
box(s,.65,1.93,12,.8,'React / TypeScript / Cytoscape: investigator interface\nPython / FastAPI / NetworkX: APIs and graph analysis; JSON + SQLite: records and decisions',18)
head(s,2.97,'Methodology and process for implementation')
labels=['Source files','Parse + extract','Review identity','Evidence graph','Signals + leads','Human review']
for i,t in enumerate(labels):
    x=.65+i*2.05
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(3.55), Inches(1.82), Inches(.63))
    sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string('EDF3F8'); sh.line.color.rgb=RGBColor.from_string(BLUE)
    box(s,x+.06,3.69,1.7,.38,t,15,True,BLUE)
    if i<5: box(s,x+1.84,3.68,.25,.35,'›',20,True,BLUE)
rows=[('ML / NLP','spaCy en_core_web_sm extracts report entities; regex extracts structured identifiers.'),('Graph algorithms','Bounded shortest paths and connected components expose recorded associations.'),('Rules + statistical ML','Temporal rules generate leads. Isolation Forest supports secondary anomaly triage.'),('Human decision','Review uncertain mentions, inspect source evidence, confirm or reject, retain audit history.')]
for i,(a,b) in enumerate(rows):
    y=4.48+i*.55; box(s,.65,y,2.7,.45,a,17,True,BLUE); box(s,3.4,y,9.2,.49,b,17)
note(s,'Separate ML, deterministic algorithms and human authority.','Local spaCy identifies entities in free text, while regex handles exact structured identifiers. Human review controls uncertain identity merges. NetworkX computes graph paths. Explicit temporal rules generate evidence-backed leads. Isolation Forest is implemented as secondary statistical triage over case features. It does not predict guilt and does not drive the primary lead-priority policy. SQLite persists review decisions.','Different methods serve distinct, auditable roles.','Why not use ML everywhere?','Exact identifiers and source-linked temporal rules benefit from deterministic behavior. Probabilistic NLP helps with unfamiliar text, but its output remains subject to review.')

s=p.slides[3]; remove(s,17410)
head(s,1.43,'Analysis of the feasibility of the idea')
box(s,.65,1.96,12,1.0,'Implemented: local frontend + backend, four input types, NER, graph traversal and evidence review.\nVerification: 98 backend tests passed on 5 September 2026 (two dependency warnings).',21)
head(s,3.13,'Potential challenges and risks')
box(s,.65,3.65,12,.93,'Synthetic case only; real-world accuracy and throughput remain unmeasured.\nNER errors, ambiguous identities, incomplete records and small anomaly cohorts can mislead.',20)
head(s,4.8,'Strategies for overcoming these challenges')
box(s,.65,5.3,12,1.33,'Current controls: source provenance, unverified extraction, reversible review and sample guards.\nProduction roadmap: agency validation, authenticated roles, encryption and protected audit storage.\nScale after profiling: asynchronous ingestion and indexed graph persistence.',19)
note(s,'Establish engineering feasibility without implying production readiness.','The current implementation runs locally and the backend test suite passed all 98 tests in this session. The main risks are data quality, extraction errors and prototype scale. Current controls include source traces and reversible human review. Production still needs authenticated access, encryption, stronger audit storage and evaluation on authorized representative data. We would profile workloads before choosing a larger graph store.','The prototype is testable, with clear production gaps.','Do 98 tests prove accuracy?','No. They verify implementation behavior and synthetic scenarios. They do not estimate operational accuracy or investigative outcomes.')

s=p.slides[4]; remove(s,17410)
head(s,1.42,'Potential impact on the target audience')
box(s,.65,1.94,12,.63,'Investigators can connect scattered records and inspect why a lead deserves review.',22,True)
box(s,.65,2.87,5.25,.45,'Benefits of the solution',20,True,BLUE)
box(s,.65,3.42,5.1,1.3,'Less repeated cross-referencing.\nSource-linked explanations.\nReviewable human decisions.',20)
box(s,.65,4.93,5.2,.4,'Synthetic demonstration',19,True,BLUE)
box(s,.65,5.4,5.15,1.18,'11 calls vs a 2/day baseline; 3 connected transfers over 75 minutes; a 6-hop path.\nTogether, these create a review-priority lead.',19)
pic(s,'04-lead-explanation.png',6.05,2.84,6.6)
box(s,6.05,6.59,6.6,.25,'Expected benefits; no measured time savings or crime-detection claim.',12)
note(s,'Connect the evidence workflow to user benefit.','In our synthetic case, a phone records eleven calls against a median baseline of two a day. Three connected transfers occur within seventy-five minutes, and a sourced six-hop path provides graph context. The backend combines these categories into a review-priority lead. The investigator can inspect the calculation and source evidence. We expect less repeated cross-referencing, but have not measured time savings.','VEIL surfaces the connection; the investigator decides what it means.','Could this be ordinary activity?','Yes. Business calls and legitimate transfers can produce the same pattern. Review priority is not a probability of crime.')

s=p.slides[5]; remove(s,17410)
head(s,1.43,'Details / Links of the reference and research work')
refs=[('spaCy English pipeline','Local named-entity extraction; en_core_web_sm','https://spacy.io/models/en#en_core_web_sm'),('NetworkX shortest paths','Graph traversal methods used by the prototype','https://networkx.org/documentation/stable/reference/algorithms/shortest_paths.html'),('scikit-learn Isolation Forest','Unsupervised anomaly detection for secondary triage','https://scikit-learn.org/stable/modules/generated/sklearn.ensemble.IsolationForest.html')]
for i,(title,desc,url) in enumerate(refs):
    y=2.06+i*1.05
    sh=box(s,.65,y,12,.35,title,20,True,BLUE); sh.text_frame.paragraphs[0].runs[0].hyperlink.address=url
    box(s,.65,y+.38,12,.3,desc,17)
    sh=box(s,.65,y+.7,12,.26,url,12); sh.text_frame.paragraphs[0].runs[0].hyperlink.address=url
head(s,5.4,'Project evidence and validation')
box(s,.65,5.89,12,.9,'Source: backend/ner_service.py, behavioral_profiler.py, demo_paths.py and lead_engine.py.\nEvidence: tests/ (98 passed), artifacts/first-round/api/ and actual prototype screenshots.\nEvaluation boundary: synthetic data; operational accuracy requires agency validation.',16)
note(s,'Make the technical foundation and evidence verifiable.','These references document the underlying NLP, graph and anomaly methods. Our contribution is their integration with evidence provenance and investigator review. The repository supplies implementation evidence, tests and captured API outputs. Our evaluation uses synthetic records, so it cannot support real-world accuracy claims. We can now demonstrate the exact route from a source record to a reviewed lead.','Claims are inspectable and limitations explicit.','Where is your evaluation dataset?','The repository contains synthetic demonstration records and test cases. Representative authorized agency data and independent validation remain future work.')

p.save(OUT/'VEIL-SIH2026.pptx')
report=['# VEIL slide-by-slide content and speaking notes','The supplied SIH template controls the six-slide structure. Its SIH images, badges, footer tags, section headings, masters and dimensions are retained. Official identity details remain placeholders.']
for i,s in enumerate(p.slides):
    report += [f'\n## Slide {i+1}', '\n'.join(sh.text for sh in s.shapes if sh.has_text_frame), notes[i]]
(OUT/'Slide-content-and-speaker-notes.md').write_text('\n\n'.join(report),encoding='utf-8')

